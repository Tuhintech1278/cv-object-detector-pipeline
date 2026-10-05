"""
Builds the complete performance whitepaper (Word .docx, ~5-6 pages) from your
measured results. All text is written for you; all numbers, tables and charts
are filled in automatically from the CSV files made by pipeline.py.

Steps:
  1. python pipeline.py --source mytest.mp4 --mode contour --no-display --max-frames 300
  2. python pipeline.py --source mytest.mp4 --mode yolo    --no-display --max-frames 300
  3. python make_report.py --author "Your Name" --org "Company Name"

Output: results/Whitepaper.docx
"""

import argparse
import ast
import glob
import os
import platform
from collections import Counter
from datetime import date

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

WARMUP = 5            # first frames ignored (model/camera warm-up)
REALTIME_MS = 1000 / 30.0
FIGDIR = os.path.join("results", "figures")
COLORS = {"contour": "#2b7bba", "yolo": "#d9822b"}
NAMES = {"contour": "Classical (OpenCV contour)", "yolo": "YOLOv8 (deep learning)"}
SHORT = {"contour": "classical OpenCV", "yolo": "YOLOv8"}


# ----------------------------------------------------------------------------
# data loading + statistics
# ----------------------------------------------------------------------------
def load_latest_per_mode():
    data = {}
    for path in sorted(glob.glob("results/metrics_*.csv")):
        df = pd.read_csv(path)
        if df.empty:
            continue
        data[df["mode"].iloc[0]] = (path, df)          # later files overwrite earlier
    return data


def stats(df):
    d = df.iloc[WARMUP:] if len(df) > WARMUP + 15 else df
    s = d["total_ms"]
    totals = Counter()
    for c in d["counts"]:
        try:
            totals.update(ast.literal_eval(c))
        except (ValueError, SyntaxError):
            pass
    return {
        "frames": len(df), "used": len(d),
        "mean": s.mean(), "median": s.median(), "p95": s.quantile(0.95),
        "max": s.max(), "std": s.std(), "fps": 1000 / s.mean(),
        "rt_pct": (s <= REALTIME_MS).mean() * 100,
        "pre": d["preprocess_ms"].mean(), "mid": d["threshold_or_post_ms"].mean(),
        "det": d["detect_ms"].mean(),
        "obj_mean": d["objects"].mean(), "obj_std": d["objects"].std(),
        "classes": totals, "d": d,
    }


# ----------------------------------------------------------------------------
# figures
# ----------------------------------------------------------------------------
def fig_architecture(path):
    fig, ax = plt.subplots(figsize=(10, 1.7))
    ax.axis("off")
    labels = ["Video source\n(webcam / file)", "Pre-processing\n(grayscale,\nGaussian blur)",
              "Detector\n(threshold + contours\nor YOLOv8)", "Overlay\n(shapes, class\ncounts)",
              "Metrics log\n(CSV per frame)"]
    n = len(labels)
    for i, text in enumerate(labels):
        cx = 0.1 + 0.2 * i
        ax.text(cx, 0.5, text, ha="center", va="center", fontsize=8, transform=ax.transAxes,
                bbox=dict(boxstyle="round,pad=0.5", fc="#e8f1fa", ec="#2b7bba"))
        if i < n - 1:
            ax.annotate("", xy=(cx + 0.185, 0.5), xytext=(cx + 0.115, 0.5), xycoords="axes fraction",
                        arrowprops=dict(arrowstyle="->", color="#444"))
    fig.savefig(path, dpi=200, bbox_inches="tight")
    plt.close(fig)


def fig_latency(st, path):
    fig, ax = plt.subplots(figsize=(7.2, 3.0))
    for mode, s in st.items():
        ax.plot(s["d"]["frame"], s["d"]["total_ms"], lw=0.9, color=COLORS[mode], label=NAMES[mode])
    ax.axhline(REALTIME_MS, ls="--", color="grey", lw=1)
    ax.text(ax.get_xlim()[1] * 0.99, REALTIME_MS, " 30 FPS limit (33.3 ms)", ha="right", va="bottom",
            fontsize=7, color="grey")
    ax.set(xlabel="Frame number", ylabel="Latency per frame (ms)")
    ax.legend(fontsize=8)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def fig_stages(st, path):
    fig, ax = plt.subplots(figsize=(7.2, 2.6))
    modes = list(st)
    parts = [("pre", "Pre-processing", "#9ec5e8"), ("mid", "Threshold / post-process", "#f2c27a"),
             ("det", "Detection / inference", "#7fbf8a")]
    left = [0] * len(modes)
    for key, lab, col in parts:
        vals = [st[m][key] for m in modes]
        ax.barh([NAMES[m] for m in modes], vals, left=left, label=lab, color=col)
        left = [l + v for l, v in zip(left, vals)]
    ax.set_xlabel("Mean time per frame (ms)")
    ax.legend(fontsize=7, loc="lower right")
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def fig_objects(st, path):
    fig, ax = plt.subplots(figsize=(7.2, 2.6))
    for mode, s in st.items():
        ax.plot(s["d"]["frame"], s["d"]["objects"], lw=0.9, color=COLORS[mode], label=NAMES[mode])
    ax.set(xlabel="Frame number", ylabel="Objects detected")
    ax.legend(fontsize=8)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


# ----------------------------------------------------------------------------
# docx helpers
# ----------------------------------------------------------------------------
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(9)


def table(doc, header, rows, widths=None, font=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font)
        shade(c, "D9E8F5")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(font)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t


def caption(doc, text):
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(8)


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(6)
    return p


def bullet(doc, text, lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if lead:
        p.add_run(lead).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(2)


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--title", default="Real-Time Object Detector and Frame-by-Frame Segmenter Pipeline: "
                                       "A Model Performance Whitepaper")
    ap.add_argument("--author", default="Your Name")
    ap.add_argument("--org", default="Your Organisation")
    ap.add_argument("--cpu", default=platform.processor() or "Laptop CPU", help="e.g. 'Intel Core i5-1135G7'")
    ap.add_argument("--ram", default="8 GB (edit)", help="e.g. '16 GB'")
    ap.add_argument("--gpu", default="None used (CPU inference)")
    ap.add_argument("--video", default="recorded test video", help="describe the input video")
    ap.add_argument("--resolution", default="640x480")
    ap.add_argument("--blur", default="5")
    ap.add_argument("--block", default="21")
    ap.add_argument("--c", default="5")
    ap.add_argument("--min-area", default="800")
    ap.add_argument("--max-area", default="100000")
    ap.add_argument("--conf", default="0.4")
    ap.add_argument("--model", default="yolov8n.pt")
    a = ap.parse_args()

    data = load_latest_per_mode()
    if not data:
        raise SystemExit("No results/metrics_*.csv files found. Run pipeline.py first.")
    st = {m: stats(df) for m, (_, df) in data.items()}
    both = "contour" in st and "yolo" in st
    if not both:
        print("[warning] only one mode found; run both modes for a full comparison. Report will still be built.")

    os.makedirs(FIGDIR, exist_ok=True)
    f_arch, f_lat = os.path.join(FIGDIR, "rep_arch.png"), os.path.join(FIGDIR, "rep_latency.png")
    f_stg, f_obj = os.path.join(FIGDIR, "rep_stages.png"), os.path.join(FIGDIR, "rep_objects.png")
    fig_architecture(f_arch)
    fig_latency(st, f_lat)
    fig_stages(st, f_stg)
    fig_objects(st, f_obj)

    # ---------------- document ----------------
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    add_page_number(sec)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)

    # title block
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(a.title)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.add_run(f"{a.author}  |  {a.org}  |  {date.today().strftime('%d %B %Y')}").font.size = Pt(10.5)

    c, y = st.get("contour"), st.get("yolo")
    ref = c or y

    # ---- Abstract
    doc.add_heading("Abstract", 1)
    if both:
        fast_m, slow_m = (("contour", "yolo") if c["mean"] < y["mean"] else ("yolo", "contour"))
        ratio = st[slow_m]["mean"] / st[fast_m]["mean"]
        abstract = (
            "This paper presents a real-time computer vision pipeline that detects and tracks target contours "
            "in video streams and overlays live per-class object counts on every frame. Two approaches are "
            "implemented and compared: a classical OpenCV pipeline (Gaussian filtering, adaptive thresholding, "
            "contour extraction with pixel-area bounds) and a pre-trained YOLOv8 deep-learning detector. Both "
            "were run on the same input and instrumented to record per-frame latency, throughput and detection "
            f"counts. The {SHORT[fast_m]} pipeline averaged {st[fast_m]['mean']:.1f} ms per frame "
            f"({st[fast_m]['fps']:.1f} FPS), while the {SHORT[slow_m]} pipeline averaged "
            f"{st[slow_m]['mean']:.1f} ms ({st[slow_m]['fps']:.1f} FPS), a factor of {ratio:.1f}. "
            "The results quantify the speed versus recognition-capability trade-off between the two approaches "
            "and provide guidance on when each is appropriate.")
    else:
        m = "contour" if c else "yolo"
        abstract = (
            "This paper presents a real-time computer vision pipeline that detects objects in video streams and "
            "overlays live per-class counts on every frame. The pipeline was instrumented to record per-frame "
            f"latency, throughput and detection counts. The {SHORT[m]} pipeline averaged "
            f"{st[m]['mean']:.1f} ms per frame ({st[m]['fps']:.1f} FPS) over {st[m]['used']} analysed frames.")
    para(doc, abstract)

    # ---- 1 Introduction
    doc.add_heading("1. Introduction", 1)
    para(doc, "Object detection is the task of locating and naming objects inside an image, and segmentation goes "
              "further by outlining each object at pixel level. Applied to video, these tasks must be repeated "
              "for every frame, so processing speed is as important as correctness: a system running at 30 frames "
              "per second has only about 33 milliseconds to analyse each frame, otherwise the output lags behind "
              "reality. Real-time detection underpins applications such as traffic monitoring, industrial "
              "inspection, retail analytics and robotics.")
    para(doc, "The objective of this project is to build a real-time pipeline that extracts and tracks target "
              "contours, overlays dynamic text counts for each detected class, and evaluates the system through "
              "frame-by-frame performance metrics. Two complementary approaches are studied. The first uses "
              "classical image-processing operations from the OpenCV library. The second uses a pre-trained "
              "deep neural network, YOLOv8. Comparing them on identical input shows what each approach costs "
              "in computation and what it offers in recognition capability.")

    # ---- 2 Architecture
    doc.add_heading("2. System Architecture", 1)
    para(doc, "The system is implemented in Python using OpenCV for video capture, image processing and display, "
              "the Ultralytics package for YOLOv8 inference, and pandas and matplotlib for analysis. Figure 1 "
              "shows the data flow. Frames are read one at a time from a webcam or video file, processed by the "
              "selected detector, annotated with the detected shapes and a live count panel, displayed, and "
              "simultaneously logged. Each frame is timed with a high-resolution clock, and the timings are "
              "written to a CSV file, which later feeds the evaluation in Section 5.")
    doc.add_picture(f_arch, width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "Figure 1. Pipeline architecture.")

    # ---- 3 Methodology
    doc.add_heading("3. Methodology", 1)
    doc.add_heading("3.1 Classical pipeline", 2)
    para(doc, "The classical detector converts each frame into a clean black-and-white mask from which object "
              "outlines are extracted. It proceeds through the following stages:")
    bullet(doc, " The colour frame is converted to a single-channel brightness image, since shape extraction "
                "does not require colour and the data volume drops to one third.", "Grayscale conversion.")
    bullet(doc, f" A Gaussian filter with a {a.blur}x{a.blur} kernel replaces each pixel with a weighted average "
                "of its neighbours, suppressing sensor noise that would otherwise create false edges.",
           "Gaussian smoothing.")
    bullet(doc, f" Instead of one global brightness cut-off, a Gaussian-weighted threshold is computed for each "
                f"{a.block}x{a.block}-pixel neighbourhood (constant C = {a.c}). This makes the method robust to "
                "uneven illumination and shadows, where a global threshold fails.", "Adaptive thresholding.")
    bullet(doc, " Morphological opening removes isolated noise specks and closing fills small gaps, producing "
                "solid blobs.", "Morphological clean-up.")
    bullet(doc, f" External contours are extracted from the mask. Only contours whose area lies between "
                f"{a.min_area} and {a.max_area} pixels are retained, which acts as a pixel-bounds filter "
                "against noise and background regions.", "Contour extraction and pixel bounds.")
    bullet(doc, " Each retained contour is simplified with the Douglas-Peucker algorithm and its corners are "
                "counted: three corners give a triangle, four a square or rectangle (decided by aspect ratio), "
                "five a pentagon, and shapes with many corners and high circularity are labelled circles.",
           "Shape classification.")
    doc.add_heading("3.2 Deep-learning pipeline", 2)
    para(doc, f"The second detector uses the pre-trained {a.model} model from the YOLO (You Only Look Once) "
              "family, trained on the COCO dataset of 80 everyday object classes. YOLO processes the entire frame "
              "in a single forward pass of a convolutional neural network and outputs bounding boxes, class "
              f"labels and confidence scores. Detections below a confidence threshold of {a.conf} are discarded. "
              "Segmentation variants of the model (suffix -seg) additionally return a pixel mask per object, "
              "enabling frame-by-frame segmentation.")
    doc.add_heading("3.3 Dynamic count overlay and metrics", 2)
    para(doc, "After detection, the objects in each frame are tallied per class and rendered as a semi-transparent "
              "text panel in the top-left corner of the output, together with the active mode, rolling FPS and "
              "frame index. The panel updates every frame, so the counts follow the scene dynamically. For each "
              "frame the system records pre-processing time, thresholding or post-processing time, detection or "
              "inference time, total latency, throughput and object count.")

    # ---- 4 Setup
    doc.add_heading("4. Experimental Setup", 1)
    para(doc, "All experiments ran on a single Windows laptop without a dedicated GPU unless stated. Both "
              "detectors processed the same input to ensure a fair comparison. The first "
              f"{WARMUP} frames of each run were excluded from the statistics because model loading and camera "
              "initialisation inflate their latency.")
    table(doc, ["Item", "Value"], [
        ["Processor", a.cpu], ["Memory", a.ram], ["GPU", a.gpu],
        ["Software", f"Python {platform.python_version()}, OpenCV, Ultralytics YOLOv8, Windows"],
        ["Input", f"{a.video}, {a.resolution} px"],
        ["Frames analysed", " / ".join(f"{NAMES[m]}: {s['used']}" for m, s in st.items())],
        ["Classical parameters", f"blur {a.blur}, block {a.block}, C {a.c}, area {a.min_area}-{a.max_area} px"],
        ["YOLO parameters", f"model {a.model}, confidence {a.conf}"],
    ], widths=[1.6, 4.9])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---- 5 Results
    doc.add_heading("5. Results", 1)
    doc.add_heading("5.1 Latency and throughput", 2)
    rows = []
    for m, s in st.items():
        rows.append([NAMES[m], f"{s['mean']:.2f}", f"{s['median']:.2f}", f"{s['p95']:.2f}",
                     f"{s['max']:.2f}", f"{s['std']:.2f}", f"{s['fps']:.1f}", f"{s['rt_pct']:.0f}%"])
    table(doc, ["Pipeline", "Mean (ms)", "Median (ms)", "P95 (ms)", "Max (ms)", "Std (ms)", "FPS", "Frames <= 33 ms"],
          rows, widths=[1.7, 0.7, 0.75, 0.65, 0.65, 0.6, 0.5, 0.85], font=8.5)
    caption(doc, "Table 1. Per-frame latency statistics (lower latency is better).")
    para(doc, "Mean latency describes typical speed, the 95th percentile (P95) shows the slow frames that "
              "viewers would notice as stutter, and the final column reports the share of frames that met the "
              "33.3 ms budget of 30 FPS real-time playback.")
    doc.add_picture(f_lat, width=Inches(6.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "Figure 2. Latency of every frame. The dashed line marks the 30 FPS budget.")

    doc.add_heading("5.2 Stage breakdown", 2)
    rows = [[NAMES[m], f"{s['pre']:.2f}", f"{s['mid']:.2f}", f"{s['det']:.2f}",
             f"{(s['mean'] - s['pre'] - s['mid'] - s['det']):.2f}"] for m, s in st.items()]
    table(doc, ["Pipeline", "Pre-processing (ms)", "Threshold / post (ms)", "Detection / inference (ms)",
                "Other: drawing, logging (ms)"], rows, widths=[1.7, 1.15, 1.2, 1.3, 1.15], font=8.5)
    caption(doc, "Table 2. Mean time spent in each stage.")
    doc.add_picture(f_stg, width=Inches(5.6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "Figure 3. Where the time goes in each pipeline.")

    doc.add_heading("5.3 Detection counts", 2)
    rows = []
    for m, s in st.items():
        top = ", ".join(f"{k}: {v}" for k, v in s["classes"].most_common(5)) or "none"
        rows.append([NAMES[m], f"{s['obj_mean']:.2f}", f"{s['obj_std']:.2f}", top])
    table(doc, ["Pipeline", "Mean objects / frame", "Std dev", "Most frequent classes (total detections)"],
          rows, widths=[1.7, 1.1, 0.7, 3.0], font=8.5)
    caption(doc, "Table 3. Object count statistics. A low standard deviation on a steady scene indicates stable detection.")
    doc.add_picture(f_obj, width=Inches(5.6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "Figure 4. Objects detected per frame.")

    # screenshots if the user saved any (press 's' while running)
    shots = sorted(glob.glob("results/shot_*.png"))[:2]
    if shots:
        doc.add_heading("5.4 Example output", 2)
        for sp in shots:
            doc.add_picture(sp, width=Inches(3.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption(doc, "Figure 5. Annotated frames with live class-count overlay.")

    # ---- 6 Discussion
    doc.add_heading("6. Discussion", 1)
    if both:
        meets = {m: st[m]["fps"] >= 30 for m in st}
        para(doc, f"The {SHORT[fast_m]} pipeline was the faster of the two, needing "
                  f"{st[fast_m]['mean']:.1f} ms per frame against {st[slow_m]['mean']:.1f} ms, a ratio of "
                  f"{ratio:.1f}. " + (
                      "Both pipelines exceeded the 30 FPS real-time threshold on this hardware."
                      if all(meets.values()) else
                      "Only the " + SHORT[[m for m in st if meets[m]][0]] + " pipeline reached the 30 FPS "
                      "real-time threshold on average."
                      if any(meets.values()) else
                      "Neither pipeline reached the 30 FPS real-time threshold on average on this hardware, "
                      "which is expected on a laptop CPU for neural network inference."),
             "Speed. ")
        mm = max(st, key=lambda k: st[k]["det"] / st[k]["mean"])
        para(doc, f"Figure 3 shows that for the {SHORT[mm]} pipeline the detection stage consumes "
                  f"{100 * st[mm]['det'] / st[mm]['mean']:.0f}% of the frame time. For the classical pipeline the "
                  "cost is spread over cheap filtering operations whose complexity grows with image resolution "
                  "but not with scene content, whereas neural network inference cost is dominated by the fixed "
                  "size of the network and benefits strongly from GPU acceleration.", "Where the time goes. ")
        para(doc, f"The 95th-percentile latencies ({c['p95']:.1f} ms classical, {y['p95']:.1f} ms YOLO) and maximum "
                  "values show occasional slow frames. Typical causes are operating-system scheduling, background "
                  "processes, thermal throttling of the laptop CPU, and frames containing many objects, which "
                  "increase drawing and post-processing work.", "Consistency. ")
    else:
        para(doc, "Only one pipeline was measured in this run, so no comparison is drawn.", "Speed. ")
    para(doc, "The classical method can only report geometric shapes that contrast with their background, and its "
              "output depends on hand-tuned parameters (block size, constant C, area bounds) that may need "
              "adjusting for each scene; shadows, textured backgrounds and touching objects can merge or split "
              "contours. YOLO recognises 80 real-world categories without tuning and handles clutter and varied "
              "lighting far better, at the price of more computation and a dependency on a pre-trained model.",
         "Capability trade-off. ")
    para(doc, "No manually labelled ground truth was available, so accuracy metrics such as precision, recall or "
              "mean average precision (mAP) were not computed; this study evaluates speed, stability and counting "
              "behaviour only. Results are specific to the tested hardware and input video, and the object counts "
              "in Table 3 should be interpreted as detections, not verified correct counts.", "Limitations. ")

    # ---- 7 Conclusion
    doc.add_heading("7. Conclusion and Future Work", 1)
    para(doc, "A working real-time pipeline was built that extracts object contours with Gaussian filtering and "
              "adaptive thresholding, alternatively detects objects with YOLOv8, overlays live per-class counts, "
              "and logs frame-by-frame inference metrics. The evaluation shows the classical approach is lighter "
              "and well suited to simple, high-contrast shapes, while the deep-learning approach offers general "
              "recognition at higher computational cost. Future work includes assigning persistent IDs to objects "
              "for true tracking across frames, labelling a test set to measure precision, recall and mAP, "
              "enabling GPU acceleration, fine-tuning a model on custom classes, and exporting to ONNX for faster "
              "deployment.")

    # ---- References
    doc.add_heading("References", 1)
    for ref_text in [
        "Bradski, G. (2000). The OpenCV Library. Dr. Dobb's Journal of Software Tools.",
        "Redmon, J., Divvala, S., Girshick, R., Farhadi, A. (2016). You Only Look Once: Unified, Real-Time "
        "Object Detection. IEEE CVPR.",
        "Jocher, G., Chaurasia, A., Qiu, J. (2023). Ultralytics YOLOv8. https://docs.ultralytics.com",
        "Lin, T.-Y. et al. (2014). Microsoft COCO: Common Objects in Context. ECCV.",
        "Suzuki, S., Abe, K. (1985). Topological structural analysis of digitized binary images by border "
        "following. Computer Vision, Graphics, and Image Processing, 30(1).",
    ]:
        p = doc.add_paragraph(ref_text, style="List Number")
        p.paragraph_format.space_after = Pt(1)
        for r in p.runs:
            r.font.size = Pt(9.5)

    out = os.path.join("results", "Whitepaper.docx")
    doc.save(out)
    print(f"[done] whitepaper saved: {out}")
    print("Open it in Word and check: your name, laptop specs, and that the page count is 5-6.")


if __name__ == "__main__":
    main()
